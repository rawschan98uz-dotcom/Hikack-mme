"""CSV and Excel parsing helpers for import endpoints."""

from __future__ import annotations

import csv
import io


def normalize_header(key: str) -> str:
    cleaned = str(key or '').strip().lower().replace(' ', '_').replace('-', '_').replace('.', '')
    for ch in '():;,/"\'#№':
        cleaned = cleaned.replace(ch, '')
    return cleaned.strip('_')


def normalize_phone(raw: str) -> str:
    s = str(raw or '').strip()
    for sep in ('/', ',', ';', '\n', '|'):
        if sep in s:
            s = s.split(sep)[0]
    return ''.join(ch for ch in s if ch.isdigit())


KNOWN_HEADER_MARKERS = {
    'имя', 'first_name', 'name', 'фио', 'fio', 'телефон', 'phone', 'tel', 'тел',
    'номер', 'номер_телефона', 'студент', 'ученик', 'student', 'фамилия', 'last_name',
    'ism', 'familiya', 'telefon',
}


def parse_csv_upload(uploaded_file) -> tuple[list[dict[str, str]], str | None]:
    if uploaded_file is None:
        return [], 'File is required'

    filename = (getattr(uploaded_file, 'name', '') or '').lower()

    if filename.endswith('.xlsx'):
        return _parse_xlsx(uploaded_file)
    elif filename.endswith('.xls'):
        return _parse_xls(uploaded_file)
    else:
        return _parse_csv(uploaded_file)


def _find_header_row_xlsx(sheet, max_scan: int = 15) -> tuple[int, list[str]]:
    # First pass: search for known column names in rows with at least 2 non-empty cells
    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        cells = [normalize_header(str(c or '')) for c in row if c is not None and str(c).strip()]
        if len(cells) >= 2 and any(any(m in c for m in KNOWN_HEADER_MARKERS) for c in cells):
            headers = [normalize_header(str(c or '')) if c is not None else '' for c in row]
            return row_idx, headers

    # Fallback: first row with at least 2 non-empty columns
    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        non_empty = [c for c in row if c is not None and str(c).strip()]
        if len(non_empty) >= 2:
            headers = [normalize_header(str(c or '')) if c is not None else '' for c in row]
            return row_idx, headers

    return 1, []


def _parse_xlsx(uploaded_file) -> tuple[list[dict[str, str]], str | None]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        sheet = wb.active
        if not sheet:
            return [], 'Excel sheet is empty'

        header_row_idx, headers = _find_header_row_xlsx(sheet)
        if not any(headers):
            return [], 'Header row is missing in Excel file'

        rows: list[dict[str, str]] = []
        for line_number, row_cells in enumerate(sheet.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
            normalized: dict[str, str] = {}
            has_value = False
            for header, val in zip(headers, row_cells):
                if not header:
                    continue
                if isinstance(val, float) and val.is_integer():
                    cell_str = str(int(val))
                elif val is not None:
                    cell_str = str(val).strip()
                else:
                    cell_str = ''
                if cell_str:
                    has_value = True
                normalized[header] = cell_str
            if has_value:
                normalized['_row'] = str(line_number)
                rows.append(normalized)

        if not rows:
            return [], 'No data rows found in Excel file'
        return rows, None
    except Exception as e:
        return [], f'Could not read Excel (.xlsx) file: {e}'


def _find_header_row_xls(sheet, max_scan: int = 15) -> tuple[int, list[str]]:
    limit = min(max_scan, sheet.nrows)
    for row_idx in range(limit):
        cells = [normalize_header(str(sheet.cell_value(row_idx, c) or '')) for c in range(sheet.ncols) if str(sheet.cell_value(row_idx, c)).strip()]
        if len(cells) >= 2 and any(any(m in c for m in KNOWN_HEADER_MARKERS) for c in cells):
            headers = [normalize_header(str(sheet.cell_value(row_idx, c) or '')) for c in range(sheet.ncols)]
            return row_idx, headers

    for row_idx in range(limit):
        non_empty = [c for c in range(sheet.ncols) if str(sheet.cell_value(row_idx, c)).strip()]
        if len(non_empty) >= 2:
            headers = [normalize_header(str(sheet.cell_value(row_idx, c) or '')) for c in range(sheet.ncols)]
            return row_idx, headers

    return 0, []


def _parse_xls(uploaded_file) -> tuple[list[dict[str, str]], str | None]:
    try:
        import xlrd
        content = uploaded_file.read()
        wb = xlrd.open_workbook(file_contents=content)
        sheet = wb.sheet_by_index(0)
        if sheet.nrows < 1:
            return [], 'Excel sheet is empty'

        header_row_idx, headers = _find_header_row_xls(sheet)
        if not any(headers):
            return [], 'Header row is missing in Excel file'

        rows: list[dict[str, str]] = []
        for line_number in range(header_row_idx + 1, sheet.nrows):
            normalized = {}
            has_value = False
            for col, header in enumerate(headers):
                if not header:
                    continue
                val = sheet.cell_value(line_number, col)
                if isinstance(val, float) and val.is_integer():
                    cell_str = str(int(val))
                elif val is not None:
                    cell_str = str(val).strip()
                else:
                    cell_str = ''
                if cell_str:
                    has_value = True
                normalized[header] = cell_str
            if has_value:
                normalized['_row'] = str(line_number + 1)
                rows.append(normalized)

        if not rows:
            return [], 'No data rows found in Excel file'
        return rows, None
    except Exception as e:
        return [], f'Could not read Excel (.xls) file: {e}'


def _parse_csv(uploaded_file) -> tuple[list[dict[str, str]], str | None]:
    try:
        raw_bytes = uploaded_file.read()
    except OSError:
        return [], 'Could not read uploaded file'

    for encoding in ('utf-8-sig', 'utf-8', 'cp1251'):
        try:
            content = raw_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            content = None
    if content is None:
        return [], 'File must be UTF-8 or Windows-1251 CSV'

    if not content.strip():
        return [], 'CSV file is empty'

    lines = content.splitlines()
    first_line = lines[0] if lines else ''
    delimiter = ';' if ';' in first_line and first_line.count(';') > first_line.count(',') else ','

    # Find header row if title comments appear on first lines
    header_idx = 0
    for idx, line in enumerate(lines[:10]):
        cleaned_line = line.lower()
        if any(marker in cleaned_line for marker in KNOWN_HEADER_MARKERS):
            header_idx = idx
            break

    csv_body = '\n'.join(lines[header_idx:])
    reader = csv.DictReader(io.StringIO(csv_body), delimiter=delimiter)
    if not reader.fieldnames:
        return [], 'CSV header row is missing'

    rows: list[dict[str, str]] = []
    for line_number, raw_row in enumerate(reader, start=header_idx + 2):
        normalized: dict[str, str] = {}
        has_value = False
        for key, value in raw_row.items():
            if key is None:
                continue
            cell = str(value or '').strip()
            if cell:
                has_value = True
            normalized[normalize_header(key)] = cell
        if has_value:
            normalized['_row'] = str(line_number)
            rows.append(normalized)

    if not rows:
        return [], 'No data rows found in CSV'

    return rows, None
